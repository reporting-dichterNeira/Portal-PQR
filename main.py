import os
import io
import uuid
import shutil
from fastapi import FastAPI, Request, HTTPException, Query, Form, UploadFile, File, Response, Depends
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from typing import Optional, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from database import (
    init_db, create_pqr, get_pqrs, get_pqr_by_id, 
    resolve_pqr, get_tipologias, get_tipologias_detail, add_tipologia, delete_tipologia,
    get_metrics, authenticate_user, create_user, get_all_users, get_user_by_id,
    update_user, delete_user, get_validator_stats, redigitar_pqr, get_pqrs_redigitacion,
    calculate_sla_business_hours
)
from models import UserLogin, UserCreate, UserUpdate, PQRResolve, TipologiaCreate, PQRRedigitar

app = FastAPI(title="Portal de Gestión PQR", version="2.2.0")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOADS_DIR = os.path.join(BASE_DIR, "uploads")
os.makedirs(UPLOADS_DIR, exist_ok=True)

# Montar carpeta de uploads para servir archivos
app.mount("/uploads", StaticFiles(directory=UPLOADS_DIR), name="uploads")

templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "templates"))

@app.on_event("startup")
def on_startup():
    init_db()

# --- Rutas de Vistas Web ---

@app.get("/", response_class=HTMLResponse)
def view_home(request: Request):
    return templates.TemplateResponse(request=request, name="login.html")

@app.get("/comercial", response_class=HTMLResponse)
def view_comercial(request: Request, email: Optional[str] = None):
    if not email:
        return RedirectResponse(url="/")
    return templates.TemplateResponse(request=request, name="comercial.html", context={"email": email.strip()})

@app.get("/validador", response_class=HTMLResponse)
def view_validador(request: Request):
    return templates.TemplateResponse(request=request, name="validador.html")

@app.get("/redigitacion", response_class=HTMLResponse)
def view_redigitacion(request: Request):
    return templates.TemplateResponse(request=request, name="redigitacion.html")

@app.get("/admin", response_class=HTMLResponse)
def view_admin(request: Request):
    return templates.TemplateResponse(request=request, name="admin.html")

# --- Autenticación ---

@app.post("/api/auth/login")
def api_login(data: UserLogin, response: Response):
    user = authenticate_user(data.email, data.password)
    if not user:
        raise HTTPException(status_code=401, detail="Correo o contraseña incorrectos")
    response.set_cookie(key="session_user", value=user["email"], httponly=True, max_age=86400)
    return user

@app.post("/api/auth/logout")
def api_logout(response: Response):
    response.delete_cookie(key="session_user")
    return {"status": "logged_out"}

# --- Panel Administrativo (Gestión de Usuarios y Tipologías) ---

@app.get("/api/admin/users")
def api_get_users():
    return get_all_users()

@app.post("/api/admin/users")
def api_create_user(user_data: UserCreate):
    created = create_user(
        nombre=user_data.nombre,
        email=user_data.email,
        password=user_data.password,
        rol=user_data.rol
    )
    if not created:
        raise HTTPException(status_code=400, detail="El correo ya se encuentra registrado.")
    return created

@app.put("/api/admin/users/{user_id}")
def api_update_user(user_id: int, user_data: UserUpdate):
    updated = update_user(
        user_id=user_id,
        nombre=user_data.nombre,
        email=user_data.email,
        rol=user_data.rol,
        activo=user_data.activo,
        password=user_data.password
    )
    if not updated:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    return updated

@app.delete("/api/admin/users/{user_id}")
def api_delete_user(user_id: int):
    if user_id == 1:
        raise HTTPException(status_code=400, detail="No es posible eliminar el superadministrador principal.")
    ok = delete_user(user_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    return {"status": "deleted"}

@app.get("/api/admin/tipologias-detail")
def api_get_tipologias_detail():
    return get_tipologias_detail()

@app.delete("/api/admin/tipologias/{tipologia_id}")
def api_delete_tipologia(tipologia_id: int):
    ok = delete_tipologia(tipologia_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Tipología no encontrada")
    return {"status": "deleted"}

# --- Catálogo de Tipologías para Validador ---

@app.get("/api/tipologias")
def api_get_tipologias():
    return get_tipologias()

@app.post("/api/tipologias")
def api_add_tipologia(data: TipologiaCreate):
    ok = add_tipologia(data.nombre)
    if not ok:
        raise HTTPException(status_code=400, detail="La tipología ya existe o no se pudo agregar.")
    return {"status": "success", "nombre": data.nombre}

# --- Estadísticas y Métricas de Rendimiento de Validadores (SLA) ---

@app.get("/api/validador/stats")
def api_get_validador_stats():
    return get_validator_stats()

# --- PQRs (Radicación Comercial con Archivos) ---

@app.post("/api/pqr")
async def api_create_pqr(
    comercial_email: str = Form(...),
    id_pdv: str = Form(...),
    cliente: str = Form(...),
    pais: str = Form(...),
    error: str = Form(...),
    archivo: Optional[UploadFile] = File(None)
):
    adjunto_nombre = None
    adjunto_ruta = None

    if archivo and archivo.filename:
        safe_name = f"{uuid.uuid4().hex[:8]}_{os.path.basename(archivo.filename)}"
        file_path = os.path.join(UPLOADS_DIR, safe_name)
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(archivo.file, buffer)
        adjunto_nombre = archivo.filename
        adjunto_ruta = f"/uploads/{safe_name}"

    created = create_pqr({
        "comercial_email": comercial_email,
        "id_pdv": id_pdv,
        "cliente": cliente,
        "pais": pais,
        "error": error,
        "adjunto_nombre": adjunto_nombre,
        "adjunto_ruta": adjunto_ruta
    })
    return created

@app.get("/api/pqr")
def api_get_pqrs(
    comercial_email: Optional[str] = None,
    estado: Optional[str] = None,
    aplica: Optional[str] = None,
    search: Optional[str] = None
):
    return get_pqrs(comercial_email=comercial_email, estado=estado, aplica=aplica, search=search)

@app.get("/api/pqr/mis-solicitudes")
def api_mis_solicitudes(email: str = Query(...)):
    return get_pqrs(comercial_email=email)

@app.get("/api/pqr/{pqr_id}")
def api_get_pqr(pqr_id: int):
    item = get_pqr_by_id(pqr_id)
    if not item:
        raise HTTPException(status_code=404, detail="PQR no encontrada")
    return item

@app.post("/api/pqr/{pqr_id}/resolver")
def api_resolve_pqr(pqr_id: int, resolve_data: PQRResolve):
    existing = get_pqr_by_id(pqr_id)
    if not existing:
        raise HTTPException(status_code=404, detail="PQR no encontrada")
    
    updated = resolve_pqr(
        pqr_id=pqr_id,
        estado=resolve_data.estado,
        aplica=resolve_data.aplica,
        tipologia=resolve_data.tipologia,
        adjudicable=resolve_data.adjudicable,
        respuesta=resolve_data.respuesta,
        validador_email=resolve_data.validador_email,
        requiere_redigitacion=resolve_data.requiere_redigitacion or False
    )
    return updated

# --- Endpoints de Redigitación ---

@app.get("/api/redigitacion")
def api_get_redigitacion(estado_redigitacion: Optional[str] = None, search: Optional[str] = None):
    return get_pqrs_redigitacion(estado_redigitacion=estado_redigitacion, search=search)

@app.post("/api/pqr/{pqr_id}/redigitar")
def api_redigitar_pqr(pqr_id: int, data: PQRRedigitar):
    existing = get_pqr_by_id(pqr_id)
    if not existing:
        raise HTTPException(status_code=404, detail="PQR no encontrada")
    if existing.get("requiere_redigitacion") != 1:
        raise HTTPException(status_code=400, detail="Esta PQR no está marcada para redigitación")
    
    updated = redigitar_pqr(
        pqr_id=pqr_id,
        nuevo_numero_auditoria=data.nuevo_numero_auditoria,
        notas=data.notas,
        redigitador_email=data.redigitador_email
    )
    return updated

@app.get("/api/metrics")
def api_get_metrics():
    return get_metrics()

# --- Exportación a Excel con Hoja de Casos y Hoja de Estadísticas de Validadores ---

@app.get("/api/pqr/export/excel")
def export_excel():
    pqrs = get_pqrs()
    stats = get_validator_stats()
    
    wb = openpyxl.Workbook()
    
    # 1. Hoja de Casos Detallados
    ws_casos = wb.active
    ws_casos.title = "Casos PQRs Detallados"
    
    headers_casos = [
        "Radicado", "Fecha Radicado", "Correo Comercial", "ID PDV", "Cliente",
        "País", "Error Reportado", "Adjunto", "Estado", "¿Aplica?", 
        "Tipología", "Área Adjudicable", "Respuesta Validador", "Fecha Resolución", "Validador",
        "Tiempo SLA (Horas Hábiles)", "Cumple SLA (≤48h)",
        "¿Requiere Redigitación?", "Estado Redigitación", "Nuevo N° Auditoría", "Fecha Redigitación", "Redigitador"
    ]
    
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    
    ws_casos.append(headers_casos)
    for col_num in range(1, len(headers_casos) + 1):
        cell = ws_casos.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for row_idx, p in enumerate(pqrs, start=2):
        f_crea = p.get("fecha_creacion")
        f_reso = p.get("fecha_resolucion")
        dur_h_str = "-"
        cumple_sla_str = "-"
        if f_crea and f_reso:
            h = calculate_sla_business_hours(f_crea, f_reso)
            dur_h_str = f"{h} h"
            cumple_sla_str = "Sí (≤48h)" if h <= 48.0 else "No (>48h)"

        row_data = [
            p.get("consecutivo", ""),
            p.get("fecha_creacion", ""),
            p.get("comercial_email", ""),
            p.get("id_pdv", ""),
            p.get("cliente", ""),
            p.get("pais", ""),
            p.get("error", ""),
            p.get("adjunto_nombre", "") or "Sin adjunto",
            p.get("estado", ""),
            p.get("aplica", "") or "Sin Dictamen",
            p.get("tipologia", "") or "Sin Asignar",
            p.get("adjudicable", "") or "Sin Asignar",
            p.get("respuesta_validador", "") or "",
            p.get("fecha_resolucion", "") or "Pendiente",
            p.get("validador_email", "") or "",
            dur_h_str,
            cumple_sla_str,
            "Sí" if p.get("requiere_redigitacion") == 1 else "No",
            p.get("estado_redigitacion", "No Aplica"),
            p.get("nuevo_numero_auditoria", "") or ("Pendiente" if p.get("requiere_redigitacion") == 1 else "N/A"),
            p.get("fecha_redigitacion", "") or "-",
            p.get("redigitador_email", "") or "-"
        ]
        ws_casos.append(row_data)
        for col_num in range(1, len(row_data) + 1):
            cell = ws_casos.cell(row=row_idx, column=col_num)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top")

    for col in ws_casos.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws_casos.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)
        
    # 2. Hoja de Estadísticas Generales y Rendimiento de Validadores (SLA)
    ws_stats = wb.create_sheet(title="Estadísticas Validadores & SLA")
    
    title_font = Font(name="Calibri", size=14, bold=True, color="1E1B4B")
    subtitle_font = Font(name="Calibri", size=11, bold=True, color="334155")
    val_header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    
    ws_stats.append(["INFORME DE RENDIMIENTO POR VALIDADOR Y MÉTRICAS SLA (≤ 48H HÁBILES)"])
    ws_stats.cell(row=1, column=1).font = title_font
    ws_stats.append(["* SLA: 48 horas hábiles excluyendo fin de semana. Casos radicados viernes > 5:00 PM inician lunes 08:00 AM."])
    ws_stats.append([])
    
    # Tabla de Validadores
    ws_stats.append(["DESGLOSE POR USUARIO VALIDADOR"])
    ws_stats.cell(row=4, column=1).font = subtitle_font
    
    val_headers = [
        "Validador", "Correo", "Tickets Resueltos", "Aplica", "% Aplica", 
        "No Aplica", "% No Aplica", "Tiempo Promedio Hábil", "Cumplimiento SLA (Meta ≤ 48h)"
    ]
    ws_stats.append(val_headers)
    header_row_idx = 5
    for col_num in range(1, len(val_headers) + 1):
        cell = ws_stats.cell(row=header_row_idx, column=col_num)
        cell.fill = val_header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    curr_row = 6
    for v in stats.get("validadores", []):
        row_vals = [
            v["nombre"],
            v["email"],
            v["total_resueltos"],
            v["aplica_count"],
            f"{v['aplica_pct']}%",
            v["no_aplica_count"],
            f"{v['no_aplica_pct']}%",
            f"{v['avg_hours']} h",
            f"{v['sla_pct']}%"
        ]
        ws_stats.append(row_vals)
        for col_num in range(1, len(row_vals) + 1):
            cell = ws_stats.cell(row=curr_row, column=col_num)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col_num > 2 else "left")
        curr_row += 1
        
    curr_row += 2
    
    # Tabla por Área Adjudicable (IT, Comercial, Campo, Validación)
    ws_stats.cell(row=curr_row, column=1, value="DISTRIBUCIÓN POR ÁREA ADJUDICABLE (IT, COMERCIAL, CAMPO, VALIDACIÓN)").font = subtitle_font
    curr_row += 1
    adj_headers = ["Área Adjudicable", "Total Casos", "Aplica", "No Aplica", "% Aplica"]
    ws_stats.append(adj_headers)
    for col_num in range(1, len(adj_headers) + 1):
        cell = ws_stats.cell(row=curr_row, column=col_num)
        cell.fill = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    curr_row += 1
    
    for a in stats.get("adjudicable_stats", []):
        row_vals = [
            a["adjudicable"],
            a["total"],
            a["aplica_cnt"],
            a["no_aplica_cnt"],
            f"{a['aplica_pct']}%"
        ]
        ws_stats.append(row_vals)
        for col_num in range(1, len(row_vals) + 1):
            cell = ws_stats.cell(row=curr_row, column=col_num)
            cell.border = thin_border
        curr_row += 1
        
    curr_row += 2

    # Tabla por Tipologías
    ws_stats.cell(row=curr_row, column=1, value="DISTRIBUCIÓN POR TIPOLOGÍA DE PQR").font = subtitle_font
    curr_row += 1
    tip_headers = ["Tipología", "Total Tickets", "Aplica", "No Aplica", "% Aplica"]
    ws_stats.append(tip_headers)
    for col_num in range(1, len(tip_headers) + 1):
        cell = ws_stats.cell(row=curr_row, column=col_num)
        cell.fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    curr_row += 1
    
    for t in stats.get("tipologias_stats", []):
        row_vals = [
            t["tipologia"],
            t["total"],
            t["aplica_cnt"],
            t["no_aplica_cnt"],
            f"{t['aplica_pct']}%"
        ]
        ws_stats.append(row_vals)
        for col_num in range(1, len(row_vals) + 1):
            cell = ws_stats.cell(row=curr_row, column=col_num)
            cell.border = thin_border
        curr_row += 1

    for col in ws_stats.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws_stats.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 50)
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = "reporte_pqrs_con_estadisticas_sla.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
